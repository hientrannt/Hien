"""
Module ghi dữ liệu vào Excel
Tự động tạo sheet cho từng loại giấy tờ
"""
import os
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from config import SHEET_NAMES, OUTPUT_DIR


class ExcelWriter:
    """Ghi dữ liệu giấy tờ vào file Excel"""

    def __init__(self, output_file=None):
        if output_file is None:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_file = os.path.join(OUTPUT_DIR, f"KetQua_{timestamp}.xlsx")

        self.output_file = output_file
        self.wb = None
        self.sheets = {}

        # Header styles
        self.header_font = Font(bold=True, color="FFFFFF")
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Border
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def _create_or_load_workbook(self):
        """Tạo hoặc load workbook"""
        if os.path.exists(self.output_file):
            self.wb = load_workbook(self.output_file)
        else:
            self.wb = Workbook()
            # Xóa sheet mặc định
            if 'Sheet' in self.wb.sheetnames:
                self.wb.remove(self.wb['Sheet'])

    def _get_or_create_sheet(self, sheet_name):
        """Lấy hoặc tạo sheet"""
        if self.wb is None:
            self._create_or_load_workbook()

        if sheet_name in self.wb.sheetnames:
            return self.wb[sheet_name]
        else:
            return self.wb.create_sheet(sheet_name)

    def _apply_header_style(self, sheet, row=1):
        """Áp dụng style cho header row"""
        for cell in sheet[row]:
            cell.font = self.header_font
            cell.fill = self.header_fill
            cell.alignment = self.header_alignment
            cell.border = self.thin_border

    def _auto_adjust_columns(self, sheet):
        """Tự động điều chỉnh độ rộng cột"""
        for column in sheet.columns:
            max_length = 0
            column_letter = column[0].column_letter

            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)
            sheet.column_dimensions[column_letter].width = adjusted_width

    def write_gpkd(self, data):
        """Ghi dữ liệu GPKD vào sheet"""
        sheet = self._get_or_create_sheet(SHEET_NAMES['GPKD'])

        # Tạo header nếu chưa có
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            headers = [
                "MST", "Số GPĐKKD", "Ngày đăng ký", "Lần thay đổi", "Ngày thay đổi",
                "Tên đơn vị", "Địa chỉ", "Điện thoại", "Email", "Vốn điều lệ",
                "Họ và tên", "Chức danh", "Giới tính", "Sinh ngày", "Quốc tịch",
                "Loại giấy tờ", "Số giấy tờ", "Ngày cấp", "Nơi cấp", "Địa chỉ thường trú",
                "Timestamp"
            ]
            sheet.append(headers)
            self._apply_header_style(sheet)

        # Thêm dữ liệu
        row_data = [
            data.get('mST', ''),
            data.get('soGPDKKD', ''),
            data.get('ngayDangKy', ''),
            data.get('lanThayDoi', ''),
            data.get('ngayThayDoi', ''),
            data.get('tenDonvi', ''),
            data.get('diaChi', ''),
            data.get('dienThoai', ''),
            data.get('emailDV', ''),
            data.get('vonDieuLe', ''),
            data.get('hoVaTen', ''),
            data.get('chucDanh', ''),
            data.get('gioiTinh', ''),
            data.get('sinhNgay', ''),
            data.get('quocTich', ''),
            data.get('loaiGiayTo', ''),
            data.get('soGiayTo', ''),
            data.get('ngayCap', ''),
            data.get('noiCap', ''),
            data.get('diaChiThuongTru', ''),
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ]
        sheet.append(row_data)

    def write_gpp(self, data):
        """Ghi dữ liệu GPP vào sheet"""
        sheet = self._get_or_create_sheet(SHEET_NAMES['GPP'])

        # Tạo header nếu chưa có
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            headers = [
                "Số GPP", "Cơ sở", "Trụ sở", "Địa chỉ",
                "Người quản lý chuyên môn", "Phạm vi kinh doanh",
                "Thời hạn", "Nơi cấp", "Ngày cấp", "Timestamp"
            ]
            sheet.append(headers)
            self._apply_header_style(sheet)

        # Thêm dữ liệu
        row_data = [
            data.get('so_gpp', ''),
            data.get('co_so', ''),
            data.get('tru_so', ''),
            data.get('dia_chi', ''),
            data.get('nguoi_quan_ly_chuyen_mon', ''),
            data.get('pham_vi_kinh_doanh', ''),
            data.get('thoi_han', ''),
            data.get('noi_cap', ''),
            data.get('ngay_cap', ''),
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ]
        sheet.append(row_data)

    def write_kdd(self, data):
        """Ghi dữ liệu KDD vào sheet"""
        sheet = self._get_or_create_sheet(SHEET_NAMES['KDD'])

        # Tạo header nếu chưa có
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            headers = [
                "Số hiệu", "Tên cơ sở", "Trụ sở", "Tên địa điểm KD", "Địa chỉ KD",
                "Họ và tên", "Trình độ chuyên môn", "Chứng chỉ hành nghề dược", "Cấp ngày",
                "Loại hình KD", "Phạm vi KD", "Thời hạn", "Nơi cấp", "Ngày cấp", "Timestamp"
            ]
            sheet.append(headers)
            self._apply_header_style(sheet)

        # Thêm dữ liệu
        row_data = [
            data.get('so_hieu', ''),
            data.get('ten_co_so', ''),
            data.get('tru_so', ''),
            data.get('ten_dia_diem_kinh_doanh', ''),
            data.get('dia_chi_kinh_doanh', ''),
            data.get('ho_va_ten', ''),
            data.get('trinh_do_chuyen_mon', ''),
            data.get('chung_chi_hanh_nghe_duoc', ''),
            data.get('cap_ngay', ''),
            data.get('du_dieu_kien_kinh_doanh_duoc_loai_hinh', ''),
            data.get('pham_vi_kinh_doanh', ''),
            data.get('thoi_han', ''),
            data.get('noi_cap', ''),
            data.get('ngay_cap', ''),
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ]
        sheet.append(row_data)

    def write_cccd(self, data):
        """Ghi dữ liệu CCCD vào sheet"""
        sheet = self._get_or_create_sheet(SHEET_NAMES['CCCD'])

        # Tạo header nếu chưa có
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            headers = [
                "Số CCCD", "Họ và tên", "Ngày sinh", "Giới tính",
                "Quốc tịch", "Quê quán", "Nơi thường trú",
                "Có giá trị đến", "Timestamp"
            ]
            sheet.append(headers)
            self._apply_header_style(sheet)

        # Thêm dữ liệu
        row_data = [
            data.get('so_cccd', ''),
            data.get('ho_va_ten', ''),
            data.get('ngay_sinh', ''),
            data.get('gioi_tinh', ''),
            data.get('quoc_tich', ''),
            data.get('que_quan', ''),
            data.get('noi_thuong_tru', ''),
            data.get('co_gia_tri_den', ''),
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ]
        sheet.append(row_data)

    def write_document(self, document):
        """
        Ghi một document vào Excel

        Args:
            document: dict {'type': 'GPKD/GPP/KDD/CCCD', 'data': {...}}

        Returns:
            dict: {'success': True/False, 'message': '...'}
        """
        try:
            doc_type = document.get('type')
            doc_data = document.get('data', {})

            if doc_type == 'GPKD':
                self.write_gpkd(doc_data)
            elif doc_type == 'GPP':
                self.write_gpp(doc_data)
            elif doc_type == 'KDD':
                self.write_kdd(doc_data)
            elif doc_type == 'CCCD':
                self.write_cccd(doc_data)
            else:
                return {
                    'success': False,
                    'message': f'Loại giấy tờ không hợp lệ: {doc_type}'
                }

            return {
                'success': True,
                'message': f'Đã ghi {doc_type} thành công'
            }

        except Exception as e:
            return {
                'success': False,
                'message': f'Lỗi ghi {doc_type}: {str(e)}'
            }

    def save(self):
        """Lưu file Excel"""
        try:
            # Auto-adjust columns
            for sheet_name in self.wb.sheetnames:
                sheet = self.wb[sheet_name]
                self._auto_adjust_columns(sheet)

            # Save file
            self.wb.save(self.output_file)
            print(f"✅ Đã lưu file: {self.output_file}")
            return True
        except Exception as e:
            print(f"❌ Lỗi lưu file: {str(e)}")
            return False

    def get_output_file(self):
        """Lấy đường dẫn file output"""
        return self.output_file


# Test function
if __name__ == "__main__":
    writer = ExcelWriter()

    # Test data
    test_gpkd = {
        'type': 'GPKD',
        'data': {
            'mST': '0123456789',
            'tenDonvi': 'CÔNG TY TEST',
            'diaChi': '123 ABC, Hà Nội'
        }
    }

    test_cccd = {
        'type': 'CCCD',
        'data': {
            'so_cccd': '001234567890',
            'ho_va_ten': 'NGUYỄN VĂN A',
            'ngay_sinh': '01/01/1990'
        }
    }

    writer.write_document(test_gpkd)
    writer.write_document(test_cccd)
    writer.save()

    print(f"File output: {writer.get_output_file()}")
